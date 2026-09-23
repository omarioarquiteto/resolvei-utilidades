from __future__ import annotations

import json
import os
import tempfile
import shutil
import zipfile
import subprocess
import re
import base64
import hashlib
import time
import threading
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
import ipaddress
import socket

import requests
from bs4 import BeautifulSoup
from dotenv import load_dotenv
try:
    from pillow_heif import register_heif_opener
    register_heif_opener()
except Exception:
    pass

from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, Response
from pydantic import BaseModel, Field

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY", "").strip()
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-5.6-luna").strip() or "gpt-5.6-luna"
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash").strip() or "gemini-3.8-flash"
SERPAPI_KEY = os.getenv("SERPAPI_KEY", "").strip()
FIREBASE_PROJECT_ID = os.getenv("FIREBASE_PROJECT_ID", "").strip()
FIREBASE_CLIENT_EMAIL = os.getenv("FIREBASE_CLIENT_EMAIL", "").strip()
FIREBASE_PRIVATE_KEY = os.getenv("FIREBASE_PRIVATE_KEY", "").replace("\\n", "\n").strip()
RESOLVEI_CREDENTIAL_ENCRYPTION_KEY = os.getenv("RESOLVEI_CREDENTIAL_ENCRYPTION_KEY", "").strip()

# Geocodificação: consultas são disparadas somente por ação explícita do usuário.
# O cache é apenas em memória para evitar chamadas repetidas ao provedor público.
_GEOCODE_CACHE: dict[str, dict[str, Any]] = {}
_CEP_CACHE: dict[str, dict[str, Any]] = {}
_NOMINATIM_LOCK = threading.Lock()
_NOMINATIM_LAST_CALL = 0.0


_firebase_admin = None
if FIREBASE_PROJECT_ID and FIREBASE_CLIENT_EMAIL and FIREBASE_PRIVATE_KEY:
    try:
        import firebase_admin
        from firebase_admin import credentials, auth as firebase_auth, firestore as firebase_firestore
        _firebase_admin = firebase_admin
        if not firebase_admin._apps:
            cred = credentials.Certificate({
                "type":"service_account","project_id":FIREBASE_PROJECT_ID,
                "private_key_id":"resolvei-env","private_key":FIREBASE_PRIVATE_KEY,
                "client_email":FIREBASE_CLIENT_EMAIL,"client_id":"",
                "token_uri":"https://oauth2.googleapis.com/token"
            })
            firebase_admin.initialize_app(cred)
    except Exception:
        _firebase_admin = None

app = FastAPI(title="Resolvei API", version="3.0.0")

# O frontend atualmente é servido pelo próprio FastAPI, portanto as requisições
# são same-origin. Mantemos CORS configurável para o domínio próprio e futuros
# clientes externos sem expor credenciais.
_default_origins = "https://resolvei-utilidades.onrender.com,https://resolvei.com.br,https://www.resolvei.com.br,http://localhost:8000,http://localhost:8080"
CORS_ORIGINS = [x.strip() for x in os.getenv("CORS_ORIGINS", _default_origins).split(",") if x.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

ALLOWED_HOSTS = {"resolvei.com.br", "www.resolvei.com.br", "localhost", "127.0.0.1"}


class RecipeRequest(BaseModel):
    url: str
    city: str = "Cuiabá"
    state: str = "MT"
    cep: str = ""
    provider: str = ""
    api_key: str = ""
    model: str = ""


class PartyRequest(BaseModel):
    type: str = "outro"
    age: int = 0
    adults: int = Field(default=20, ge=0)
    kids: int = Field(default=10, ge=0)
    hours: float = Field(default=4, ge=0.5)
    drinkers: int = Field(default=8, ge=0)
    alcohol: str = "beer"
    budgetPerPerson: float = Field(default=0, ge=0)


class ShoppingPriceRequest(BaseModel):
    item: str
    city: str = "Cuiabá"
    state: str = "MT"


class SalesPriceRequest(BaseModel):
    product: str
    type: str = "outro"
    city: str = "Cuiabá"
    state: str = "MT"
    quantity: float = Field(default=10, gt=0)
    production_cost: float = Field(default=0, ge=0)
    packaging_per_unit: float = Field(default=0, ge=0)
    other_costs: float = Field(default=0, ge=0)
    provider: str = ""
    api_key: str = ""
    model: str = ""


class SolarResourceRequest(BaseModel):
    cep: str = ""
    street: str = ""
    neighborhood: str = ""
    city: str = "Cuiabá"
    state: str = "MT"
    address: str = ""
    roofAzimuth: float = 0
    roofTilt: float = 15


class AIConnectionRequest(BaseModel):
    provider: str
    api_key: str
    model: str = ""

class AIChatRequest(BaseModel):
    message: str
    provider: str = ""
    model: str = ""

class RiskProgressRequest(BaseModel):
    state: dict[str, Any] = Field(default_factory=dict)


class SolarImageRequest(BaseModel):
    imageData: str
    terrainWidth: float = 0
    terrainDepth: float = 0
    roofAzimuth: float = 0
    lat: float = 0
    lon: float = 0


def safe_url(raw: str) -> str:
    value = raw.strip()
    if not value.startswith(("https://", "http://")):
        raise HTTPException(status_code=400, detail="Informe uma URL começando com http:// ou https://.")
    parsed = urlparse(value)
    if not parsed.netloc:
        raise HTTPException(status_code=400, detail="URL inválida.")
    host = (parsed.hostname or "").lower()
    if host in {"localhost", "127.0.0.1", "0.0.0.0", "::1"} or host.endswith(".local"):
        raise HTTPException(status_code=400, detail="Por segurança, endereços locais não podem ser consultados.")
    try:
        infos = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
        for info in infos:
            addr = info[4][0]
            ip = ipaddress.ip_address(addr)
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                raise HTTPException(status_code=400, detail="O endereço informado resolve para uma rede privada e foi bloqueado por segurança.")
    except socket.gaierror:
        raise HTTPException(status_code=400, detail="Não foi possível resolver o domínio informado.")
    return value


def fetch_html(url: str) -> str:
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; ResolveiBot/2.1; +https://resolvei.com.br/)",
        "Accept-Language": "pt-BR,pt;q=0.9,en;q=0.6",
    }
    try:
        response = requests.get(url, headers=headers, timeout=15, allow_redirects=True)
        response.raise_for_status()
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Não foi possível ler a página da receita: {exc}") from exc
    content_type = response.headers.get("content-type", "")
    if "text/html" not in content_type and "application/xhtml+xml" not in content_type:
        raise HTTPException(status_code=415, detail="A URL não parece apontar para uma página HTML.")
    return response.text[:5_000_000]


def walk_recipe_json(data: Any) -> list[str]:
    found: list[str] = []
    stack = [data]
    while stack:
        item = stack.pop()
        if isinstance(item, dict):
            typ = item.get("@type")
            types = typ if isinstance(typ, list) else [typ]
            if any(isinstance(t, str) and t.lower() == "recipe" for t in types):
                ing = item.get("recipeIngredient")
                if isinstance(ing, list):
                    found.extend(str(x).strip() for x in ing if str(x).strip())
                elif isinstance(ing, str) and ing.strip():
                    found.append(ing.strip())
            for value in item.values():
                if isinstance(value, (dict, list)):
                    stack.append(value)
        elif isinstance(item, list):
            stack.extend(item)
    return found


def extract_raw_ingredients(html: str) -> list[str]:
    soup = BeautifulSoup(html, "html.parser")
    found: list[str] = []

    for script in soup.find_all("script", attrs={"type": re.compile(r"ld\+json", re.I)}):
        raw = script.string or script.get_text(strip=True)
        if not raw:
            continue
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            continue
        found.extend(walk_recipe_json(data))

    for el in soup.find_all(attrs={"itemprop": "recipeIngredient"}):
        txt = el.get("content") or el.get_text(" ", strip=True)
        if txt:
            found.append(txt.strip())

    # Fallback: look for lists/headings with ingredient-like class names.
    if len(found) < 2:
        for el in soup.find_all(class_=re.compile(r"ingred|ingredient|ingredients|ingredientes", re.I)):
            txt = el.get_text(" ", strip=True)
            if 3 <= len(txt) <= 240 and txt.lower() not in {x.lower() for x in found}:
                found.append(txt)

    # Clean, deduplicate and avoid obvious headings.
    clean: list[str] = []
    bad = {"ingredientes", "ingredients", "modo de preparo", "modo de preparo:", "ingrediente"}
    for item in found:
        item = re.sub(r"\s+", " ", item).strip(" •-\t\r\n")
        if not item or item.lower() in bad or len(item) > 250:
            continue
        if item.lower() not in {x.lower() for x in clean}:
            clean.append(item)
    return clean[:30]


def heuristic_category(name: str) -> str:
    text = name.lower()
    groups = {
        "laticínios": ["leite", "manteiga", "queijo", "creme de leite", "iogurte", "requeijão"],
        "carnes": ["carne", "picanha", "fraldinha", "frango", "linguiça", "bacon", "presunto", "peixe"],
        "frutas": ["banana", "maçã", "laranja", "limão", "morango", "abacaxi", "mamão", "uva"],
        "hortaliças": ["tomate", "cebola", "alho", "batata", "cenoura", "pimentão", "mandioca", "cheiro-verde"],
        "secos": ["farinha", "arroz", "feijão", "açúcar", "sal", "amido", "fermento", "aveia", "macarrão"],
        "bebidas": ["água", "suco", "refrigerante", "vinho", "cerveja", "café", "leite"],
        "óleos e molhos": ["óleo", "azeite", "vinagre", "molho", "maionese"],
        "temperos": ["páprica", "pimenta", "orégano", "cominho", "canela", "cravo", "ervas"],
    }
    for cat, words in groups.items():
        if any(word in text for word in words):
            return cat
    return "outros"


def heuristic_parse(raw: str) -> dict[str, Any]:
    m = re.match(r"^\s*(\d+(?:[\.,]\d+)?)\s*([A-Za-zÀ-ÿ\.]+)?\s*(.*)$", raw)
    quantity = 1.0
    unit = "un."
    name = raw.strip()
    if m:
        quantity = float(m.group(1).replace(",", "."))
        unit_raw = (m.group(2) or "").lower().rstrip(".")
        aliases = {"g": "g", "kg": "kg", "ml": "ml", "l": "L", "xícara": "xícara", "xícaras": "xícara", "colher": "colher", "colheres": "colher", "unidade": "un.", "unidades": "un.", "dente": "un.", "dentes": "un."}
        unit = aliases.get(unit_raw, unit or "un.")
        name = m.group(3).strip(" -:") or raw.strip()
    return {"name": name, "quantity": quantity, "unit": unit, "category": heuristic_category(name), "raw": raw, "recipe_price": 0}


def _verify_firebase_token(authorization: str | None):
    if not _firebase_admin:
        raise HTTPException(status_code=503, detail="Firebase Admin não está configurado no servidor.")
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Faça login no Resolvei.")
    try:
        from firebase_admin import auth as firebase_auth
        return firebase_auth.verify_id_token(authorization[7:].strip())
    except Exception as exc:
        raise HTTPException(status_code=401, detail="Sessão Firebase inválida ou expirada.") from exc

def _credential_key() -> bytes:
    if not RESOLVEI_CREDENTIAL_ENCRYPTION_KEY:
        raise HTTPException(status_code=503, detail="RESOLVEI_CREDENTIAL_ENCRYPTION_KEY não configurada.")
    try:
        raw=base64.urlsafe_b64decode(RESOLVEI_CREDENTIAL_ENCRYPTION_KEY + "="*((4-len(RESOLVEI_CREDENTIAL_ENCRYPTION_KEY)%4)%4))
        if len(raw)!=32: raise ValueError
        return raw
    except Exception as exc:
        raise HTTPException(status_code=503, detail="A chave de criptografia deve ser Base64URL de 32 bytes.") from exc

def _encrypt_secret(secret: str) -> str:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    nonce=os.urandom(12); ct=AESGCM(_credential_key()).encrypt(nonce,secret.encode(),None)
    return base64.urlsafe_b64encode(nonce+ct).decode()

def _decrypt_secret(value: str) -> str:
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    raw=base64.urlsafe_b64decode(value.encode())
    return AESGCM(_credential_key()).decrypt(raw[:12],raw[12:],None).decode()

def _resolve_ai_credentials(authorization: str | None, provider: str = ""):
    token=_verify_firebase_token(authorization)
    from firebase_admin import firestore as firebase_firestore
    docs=firebase_firestore.client().collection("users").document(token["uid"]).collection("aiConnections").stream()
    conns={d.id:d.to_dict() or {} for d in docs}
    wanted=(provider or "").lower().strip()
    if wanted:
        conn=conns.get(wanted,{})
        if conn.get("secret"):
            return wanted,_decrypt_secret(conn["secret"]),str(conn.get("model") or "")
        if wanted=="gemini" and GEMINI_API_KEY:
            return "gemini",GEMINI_API_KEY,GEMINI_MODEL
        raise HTTPException(status_code=400,detail=f"Nenhuma conexão de IA para {wanted}.")
    for candidate in ("gemini","openai","anthropic","openrouter"):
        conn=conns.get(candidate,{})
        if conn.get("secret"):
            return candidate,_decrypt_secret(conn["secret"]),str(conn.get("model") or "")
    if GEMINI_API_KEY:
        return "gemini",GEMINI_API_KEY,GEMINI_MODEL
    raise HTTPException(status_code=400,detail="Nenhuma IA conectada e o Gemini do Resolvei não está configurado.")

def _gemini_error_detail(resp) -> str:
    try:
        data=resp.json()
        err=data.get("error",{}) if isinstance(data,dict) else {}
        return f"{err.get('code',resp.status_code)} {err.get('status','')}: {err.get('message',resp.text[:300])}"
    except Exception:
        return f"{resp.status_code}: {resp.text[:300]}"

def _gemini_available_models(api_key: str, preferred: str = "") -> list[str]:
    resp=requests.get("https://generativelanguage.googleapis.com/v1beta/models",headers={"x-goog-api-key":api_key},params={"pageSize":1000},timeout=30)
    resp.raise_for_status()
    data=resp.json()
    models=data.get("models",[]) if isinstance(data,dict) else []
    available=[]
    for item in models:
        name=str(item.get("name","")).strip()
        methods=item.get("supportedGenerationMethods",[]) or []
        if name.startswith("models/"): name=name[7:]
        if name and "generateContent" in methods: available.append(name)
    priority=[preferred.strip(),"gemini-3.8-flash","gemini-3.5-flash","gemini-3.1-flash-lite"]
    return list(dict.fromkeys([x for x in priority+available if x and x in available]))

def _gemini_generate(api_key: str, model: str, message: str) -> str:
    url="https://generativelanguage.googleapis.com/v1beta/models/"+model+":generateContent"
    payload={"contents":[{"parts":[{"text":message}]}],"generationConfig":{"responseMimeType":"application/json"}}
    last=None
    for attempt in range(3):
        resp=requests.post(url,headers={"x-goog-api-key":api_key,"Content-Type":"application/json"},json=payload,timeout=60)
        if resp.ok:
            data=resp.json()
            return data["candidates"][0]["content"]["parts"][0]["text"]
        last=RuntimeError("Gemini "+model+": "+_gemini_error_detail(resp))
        if resp.status_code in {429,500,502,503,504} and attempt<2:
            time.sleep(1.5*(attempt+1)); continue
        raise last
    raise last or RuntimeError("Gemini não retornou resposta.")

def _provider_call(provider: str, api_key: str, model: str, message: str) -> str:
    provider=provider.lower()
    if provider=="openai":
        from openai import OpenAI
        client=OpenAI(api_key=api_key)
        resp=client.responses.create(model=model or "gpt-4.1-mini",input=message)
        return getattr(resp,"output_text","") or ""
    if provider=="gemini":
        available = _gemini_available_models(api_key, model)
        if not available:
            raise RuntimeError("A chave Gemini não possui nenhum modelo com generateContent disponível.")
        errors=[]
        for candidate in available[:8]:
            try:
                return _gemini_generate(api_key, candidate, message)
            except Exception as exc:
                errors.append(str(exc))
        raise RuntimeError("Nenhum modelo Gemini disponível respondeu. " + " | ".join(errors[:4]))

    if provider=="anthropic":
        model=model or "claude-3-5-haiku-latest"
        resp=requests.post("https://api.anthropic.com/v1/messages",headers={"x-api-key":api_key,"anthropic-version":"2023-06-01","content-type":"application/json"},json={"model":model,"max_tokens":1200,"messages":[{"role":"user","content":message}]},timeout=60)
        resp.raise_for_status()
        return "".join(x.get("text","") for x in resp.json().get("content",[]))
    if provider=="openrouter":
        resp=requests.post("https://openrouter.ai/api/v1/chat/completions",headers={"Authorization":"Bearer "+api_key,"Content-Type":"application/json"},json={"model":model or "openai/gpt-4.1-mini","messages":[{"role":"user","content":message}]},timeout=60)
        resp.raise_for_status()
        return resp.json()["choices"][0]["message"]["content"]
    raise HTTPException(status_code=400,detail="Provedor de IA não suportado.")

@app.get("/api/ai/diagnose")
def ai_diagnose(authorization: str | None = Header(default=None)) -> dict[str, Any]:
    _verify_firebase_token(authorization)
    provider, api_key, stored_model = _resolve_ai_credentials(authorization, "gemini")
    result = {"provider": provider, "configured_model": stored_model or GEMINI_MODEL, "models": [], "test": None}
    try:
        models = _gemini_available_models(api_key, stored_model or GEMINI_MODEL)
        result["models"] = models[:12]
        if not models:
            result["test"] = {"ok": False, "error": "Nenhum modelo com generateContent disponível para esta chave."}
            return result
        test_errors = []
        for candidate in models[:5]:
            try:
                answer = _gemini_generate(api_key, candidate, "Responda somente: OK")
                result["test"] = {"ok": True, "model": candidate, "answer": answer[:50]}
                return result
            except Exception as exc:
                test_errors.append(str(exc))
        result["test"] = {"ok": False, "error": " | ".join(test_errors[:4])}
        return result
    except Exception as exc:
        result["test"] = {"ok": False, "error": str(exc)}
        return result

@app.get("/api/ai/connections")
def ai_connections(authorization: str | None = Header(default=None)):
    user=_verify_firebase_token(authorization)
    if not _firebase_admin: raise HTTPException(status_code=503,detail="Firebase indisponível.")
    from firebase_admin import firestore as firebase_firestore
    docs=firebase_firestore.client().collection("users").document(user["uid"]).collection("aiConnections").stream()
    return {"connections":[{"provider":d.id,"connected":bool((d.to_dict() or {}).get("connected")),"model":(d.to_dict() or {}).get("model","")} for d in docs]}

@app.post("/api/ai/connections")
def ai_save_connection(req: AIConnectionRequest, authorization: str | None = Header(default=None)):
    user=_verify_firebase_token(authorization)
    if req.provider not in {"gemini","openai","anthropic","openrouter"}: raise HTTPException(status_code=400,detail="Provedor não suportado.")
    if not req.api_key.strip(): raise HTTPException(status_code=400,detail="Informe a API Key.")
    from firebase_admin import firestore as firebase_firestore
    firebase_firestore.client().collection("users").document(user["uid"]).collection("aiConnections").document(req.provider).set({
        "provider":req.provider,"connected":True,"model":req.model.strip(),"secret":_encrypt_secret(req.api_key.strip()),
        "updatedAt":firebase_firestore.SERVER_TIMESTAMP
    },merge=True)
    return {"ok":True,"provider":req.provider,"connected":True,"model":req.model.strip()}

@app.delete("/api/ai/connections/{provider}")
def ai_delete_connection(provider: str, authorization: str | None = Header(default=None)):
    user=_verify_firebase_token(authorization)
    from firebase_admin import firestore as firebase_firestore
    firebase_firestore.client().collection("users").document(user["uid"]).collection("aiConnections").document(provider).delete()
    return {"ok":True}

@app.post("/api/ai/chat")
def ai_chat(req: AIChatRequest, authorization: str | None = Header(default=None)):
    user=_verify_firebase_token(authorization)
    from firebase_admin import firestore as firebase_firestore
    docs=firebase_firestore.client().collection("users").document(user["uid"]).collection("aiConnections").stream()
    conns={d.id:d.to_dict() or {} for d in docs}
    provider=req.provider.lower().strip()
    key=""
    model=req.model.strip()
    if provider and conns.get(provider,{}).get("secret"):
        key=_decrypt_secret(conns[provider]["secret"])
        model=model or conns[provider].get("model","")
    elif not provider:
        provider,key,stored_model=_resolve_ai_credentials(authorization,"")
        model=model or stored_model
    elif provider=="gemini" and GEMINI_API_KEY:
        key=GEMINI_API_KEY
        model=model or GEMINI_MODEL
    else:
        raise HTTPException(status_code=400,detail="Nenhuma conexão de IA disponível para este provedor.")
    try:
        answer=_provider_call(provider,key,model,req.message)
        return {"provider":provider,"model":model,"answer":answer}
    except HTTPException: raise
    except Exception as exc:
        raise HTTPException(status_code=502,detail=f"Falha ao consultar {provider}: {exc}") from exc

class AIPublicChatRequest(BaseModel):
    provider: str
    api_key: str
    message: str
    model: str = ""

@app.post("/api/ai/chat-public")
def ai_chat_public(req: AIPublicChatRequest):
    provider=req.provider.lower().strip()
    if provider not in {"gemini","openai","anthropic","openrouter"}:
        raise HTTPException(status_code=400,detail="Provedor de IA não suportado.")
    api_key=req.api_key.strip()
    message=req.message.strip()
    if not api_key:
        raise HTTPException(status_code=400,detail="Informe sua API Key.")
    if not message:
        raise HTTPException(status_code=400,detail="Informe o que você quer resolver.")
    # A chave é recebida apenas para esta requisição e não é persistida.
    try:
        answer=_provider_call(provider,api_key,req.model.strip(),message)
        return {"provider":provider,"model":req.model.strip(),"answer":answer}
    except HTTPException:
        raise
    except Exception as exc:
        detail=str(exc)
        if len(detail)>300: detail=detail[:300]
        raise HTTPException(status_code=502,detail=f"Falha ao consultar {provider}: {detail}") from exc

def openai_json(prompt: str, schema_name: str, schema: dict[str, Any]) -> dict[str, Any]:
    if not OPENAI_API_KEY:
        raise RuntimeError("OPENAI_API_KEY não configurada")
    try:
        from openai import OpenAI
        client = OpenAI(api_key=OPENAI_API_KEY)
        response = client.responses.create(
            model=OPENAI_MODEL,
            input=prompt,
            text={"format": {"type": "json_schema", "name": schema_name, "strict": True, "schema": schema}},
        )
        text = getattr(response, "output_text", "") or ""
        return json.loads(text)
    except Exception as exc:
        raise RuntimeError(f"Falha ao consultar a IA: {exc}") from exc


def normalize_ingredients(raw_items: list[str], page_title: str = "") -> list[dict[str, Any]]:
    if not raw_items:
        return []
    schema = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "ingredients": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {
                        "name": {"type": "string"},
                        "quantity": {"type": "number"},
                        "unit": {"type": "string"},
                        "category": {"type": "string"},
                        "raw": {"type": "string"},
                        "recipe_price": {"type": "number"},
                    },
                    "required": ["name", "quantity", "unit", "category", "raw", "recipe_price"],
                },
            }
        },
        "required": ["ingredients"],
    }
    prompt = f"""Normalize uma lista de ingredientes de uma receita brasileira. Preserve a quantidade original quando possível. Separe quantidade, unidade e nome do ingrediente; classifique em categorias úteis para compras. Nunca invente preços: recipe_price deve ser 0. Página: {page_title}\n\nIngredientes brutos:\n""" + "\n".join(f"- {x}" for x in raw_items)
    try:
        data = openai_json(prompt, "resolvei_recipe", schema)
        return data.get("ingredients", [])[:30]
    except RuntimeError:
        return [heuristic_parse(x) for x in raw_items]


def price_to_float(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    nums = re.findall(r"\d+[\d\.,]*", text)
    if not nums:
        return None
    candidate = nums[-1].replace(".", "").replace(",", ".")
    try:
        return float(candidate)
    except ValueError:
        return None


def search_price(item: str, city: str, state: str) -> list[dict[str, Any]]:
    if not SERPAPI_KEY:
        return []
    params = {
        "engine": "google_shopping",
        "q": item,
        "location": f"{city}, {state}, Brazil",
        "hl": "pt-BR",
        "gl": "br",
        "api_key": SERPAPI_KEY,
        "num": 6,
    }
    try:
        res = requests.get("https://serpapi.com/search.json", params=params, timeout=15)
        res.raise_for_status()
        data = res.json()
    except requests.RequestException:
        return []
    options: list[dict[str, Any]] = []
    for result in data.get("shopping_results", [])[:6]:
        price = price_to_float(result.get("price"))
        if price is None:
            continue
        options.append({
            "title": result.get("title") or item,
            "price": price,
            "source": result.get("source") or "Loja online",
            "link": result.get("link") or result.get("product_link") or "",
        })
    return options



SOLAR_MONTHS = ["Jan","Fev","Mar","Abr","Mai","Jun","Jul","Ago","Set","Out","Nov","Dez"]

def _lookup_cep(cep: str) -> dict[str, Any]:
    digits = re.sub(r"\D", "", cep or "")
    if len(digits) != 8:
        raise HTTPException(status_code=400, detail="CEP deve conter 8 dígitos.")
    if digits in _CEP_CACHE:
        return _CEP_CACHE[digits]
    try:
        r = requests.get(
            f"https://viacep.com.br/ws/{digits}/json/",
            headers={"User-Agent": "Resolvei/3.2 (address lookup)"},
            timeout=10,
        )
        r.raise_for_status()
        data = r.json()
    except requests.RequestException as exc:
        raise HTTPException(status_code=502, detail=f"Falha ao consultar CEP: {exc}") from exc
    if data.get("erro"):
        raise HTTPException(status_code=404, detail="CEP não encontrado.")
    _CEP_CACHE[digits] = data
    return data


def _nominatim_search(query: str) -> dict[str, Any]:
    global _NOMINATIM_LAST_CALL
    key = re.sub(r"\s+", " ", query.strip().lower())
    if key in _GEOCODE_CACHE:
        return _GEOCODE_CACHE[key]

    # A instância pública do Nominatim pede no máximo 1 requisição/s por aplicação.
    with _NOMINATIM_LOCK:
        wait = 1.0 - (time.monotonic() - _NOMINATIM_LAST_CALL)
        if wait > 0:
            time.sleep(wait)
        try:
            r = requests.get(
                "https://nominatim.openstreetmap.org/search",
                params={
                    "q": query,
                    "format": "json",
                    "limit": 1,
                    "countrycodes": "br",
                    "addressdetails": 1,
                },
                headers={
                    "User-Agent": "Resolvei/3.2 (solar-study; https://resolvei.com.br/)",
                    "Accept-Language": "pt-BR,pt;q=0.9",
                },
                timeout=15,
            )
            _NOMINATIM_LAST_CALL = time.monotonic()
            r.raise_for_status()
            data = r.json()
        except requests.RequestException as exc:
            _NOMINATIM_LAST_CALL = time.monotonic()
            raise HTTPException(status_code=502, detail=f"Falha no serviço de localização: {exc}") from exc
    if not data:
        raise HTTPException(
            status_code=404,
            detail="Não encontrei coordenadas para esse endereço. Tente informar rua, número, cidade e UF ou use o CEP.",
        )
    x = data[0]
    result = {
        "lat": float(x["lat"]),
        "lon": float(x["lon"]),
        "displayName": x.get("display_name", query),
        "addressDetails": x.get("address", {}),
        "geocoder": "OpenStreetMap/Nominatim",
    }
    if len(_GEOCODE_CACHE) >= 256:
        _GEOCODE_CACHE.pop(next(iter(_GEOCODE_CACHE)))
    _GEOCODE_CACHE[key] = result
    return result


def geocode_solar(req: SolarResourceRequest) -> dict[str, Any]:
    cep_data: dict[str, Any] | None = None
    if req.cep:
        cep_digits = re.sub(r"\D", "", req.cep)
        if len(cep_digits) != 8:
            raise HTTPException(status_code=400, detail="CEP deve conter 8 dígitos.")
        cep_data = _lookup_cep(cep_digits)

    street = (req.street or "").strip() or ((cep_data or {}).get("logradouro") or "").strip()
    neighborhood = (req.neighborhood or "").strip() or ((cep_data or {}).get("bairro") or "").strip()
    city = (req.city or "").strip() or ((cep_data or {}).get("localidade") or "").strip()
    state = (req.state or "").strip() or ((cep_data or {}).get("uf") or "").strip()
    address = (req.address or "").strip() or street

    pieces = [x for x in [address, neighborhood, city, state, "Brasil"] if x]
    # O CEP ajuda bastante quando o endereço é pouco específico, mas não é enviado
    # como único termo porque o Nominatim pode ter cobertura postal irregular.
    if req.cep:
        pieces.insert(-1, f"CEP {re.sub(r'\D', '', req.cep)}")
    query = ", ".join(pieces)
    if not query:
        raise HTTPException(status_code=400, detail="Informe um endereço ou CEP para localizar o ponto.")
    result = _nominatim_search(query)
    result["input"] = {
        "address": address,
        "street": street,
        "neighborhood": neighborhood,
        "city": city,
        "state": state,
        "cep": req.cep or "",
    }
    return result

def fallback_solar(lat: float, roof_azimuth: float = 0, roof_tilt: float = 15) -> dict[str, Any]:
    lat_abs = abs(lat)
    psh = min(6.2, max(3.5, 4.7 + max(0.0, 22-lat_abs)*0.025))
    preferred = 0.0 if lat < 0 else 180.0
    tilt = max(10.0, min(35.0, lat_abs))
    diff = abs(((roof_azimuth-preferred+180)%360)-180)
    orient = max(0.78, 1 - (diff/180)*0.24)
    tilt_factor = max(0.90, 1 - abs(roof_tilt-tilt)/90*0.10)
    annual = psh*365*0.80*orient*tilt_factor
    monthly_avg = annual/12
    seasonal = [1.06,1.05,1.00,0.95,0.88,0.82,0.83,0.88,0.96,1.02,1.05,1.08]
    monthly=[{"name":m,"kwh":monthly_avg*f} for m,f in zip(SOLAR_MONTHS,seasonal)]
    return {"source":"fallback","optimalAzimuth":preferred,"optimalTilt":round(tilt,1),"annualKwhPerKwp":round(annual,1),"specificMonthly":round(monthly_avg,1),"monthly":monthly,"psh":round(psh,2)}


def pvgis_fixed(lat: float, lon: float, angle: float, aspect: float) -> dict[str, Any] | None:
    params={"lat":lat,"lon":lon,"peakpower":1,"loss":14,"pvtechchoice":"crystSi","mountingplace":"building","angle":angle,"aspect":aspect,"pvcalculation":1,"outputformat":"json"}
    try:
        r=requests.get("https://re.jrc.ec.europa.eu/api/v5_3/PVcalc",params=params,timeout=20)
        r.raise_for_status(); return r.json()
    except requests.RequestException:
        return None


def pvgis_optimal(lat: float, lon: float) -> dict[str, Any] | None:
    params={"lat":lat,"lon":lon,"peakpower":1,"loss":14,"pvtechchoice":"crystSi","mountingplace":"building","optimalangles":1,"pvcalculation":1,"outputformat":"json"}
    try:
        r=requests.get("https://re.jrc.ec.europa.eu/api/v5_3/PVcalc",params=params,timeout=20)
        r.raise_for_status(); return r.json()
    except requests.RequestException:
        return None


def parse_pvgis(data: dict[str, Any]) -> dict[str, Any]:
    out=data.get("outputs",{}) or {}
    fixed=out.get("fixed",{}) or {}
    totals=out.get("totals",{}) or {}
    monthly=out.get("monthly",[]) or []
    annual = totals.get("fixed",{}).get("E_y") if isinstance(totals.get("fixed"),dict) else None
    if annual is None:
        annual = fixed.get("E_y")
    if annual is None and monthly:
        annual = sum(float(m.get("E_m",0) or 0) for m in monthly)
    rows=[]
    for i,m in enumerate(monthly[:12]):
        rows.append({"name":SOLAR_MONTHS[i] if i<len(SOLAR_MONTHS) else str(i+1),"kwh":float(m.get("E_m",0) or 0)})
    inp=data.get("inputs",{}) or {}
    fixed_in=inp.get("fixed",{}) if isinstance(inp.get("fixed"),dict) else {}
    slope=fixed_in.get("slope", data.get("meta",{}).get("slope",0))
    aspect=fixed_in.get("azimuth", fixed_in.get("aspect",0))
    return {"annualKwhPerKwp":float(annual or 0),"specificMonthly":float((annual or 0)/12),"monthly":rows,"slope":float(slope or 0),"aspect":float(aspect or 0)}


def pvgis_summary(lat: float, lon: float, roof_azimuth: float, roof_tilt: float) -> dict[str, Any]:
    optimal=pvgis_optimal(lat,lon)
    fixed_aspect=((roof_azimuth-180+180)%360)-180
    fixed=pvgis_fixed(lat,lon,roof_tilt,fixed_aspect)
    if optimal:
        opt=parse_pvgis(optimal)
        pv_aspect=opt.get("aspect",0)
        preferred=((pv_aspect+180)%360)
        return {"source":"PVGIS","optimalAzimuth":preferred,"optimalTilt":opt.get("slope",abs(lat)),"annualKwhPerKwp":opt.get("annualKwhPerKwp",0),"specificMonthly":opt.get("specificMonthly",0),"monthly":opt.get("monthly",[]),"roof":parse_pvgis(fixed) if fixed else None}
    return fallback_solar(lat,roof_azimuth,roof_tilt)

def fallback_party(req: PartyRequest) -> dict[str, Any]:
    labels = {
        "aniversario-infantil": "Aniversário infantil", "aniversario-adulto": "Aniversário adulto", "casamento": "Casamento",
        "firma": "Festa da firma / confraternização", "cha-bebe": "Chá de bebê", "cha-revelacao": "Chá revelação",
        "noivado": "Noivado", "formatura": "Formatura", "bodas": "Bodas", "familiar": "Comemoração familiar",
        "junina": "Festa junina", "happy-hour": "Happy hour", "outro": "Outra comemoração",
    }
    g=req.adults+req.kids
    drinkers=min(g,max(0,req.drinkers))
    if req.alcohol == "beer": alcohol = f"≈ {drinkers*(req.hours/4)*1.5:.1f} L de cerveja"
    elif req.alcohol == "wine": alcohol = f"≈ {max(0,round(drinkers*(req.hours/4)*0.25+0.49))} garrafas de vinho de 750 ml"
    elif req.alcohol == "drinks": alcohol = f"≈ {max(0,round(drinkers*(req.hours/4)*4))} doses de drinks"
    elif req.alcohol == "mixed": alcohol = "Mix de cerveja, vinho e/ou drinks a dimensionar pelo perfil"
    else: alcohol = "Sem álcool"
    if req.type == "junina": foods=["milho/pamonha","canjica ou curau","paçoca e pé-de-moleque","cachorro-quente"]
    elif req.type == "happy-hour": foods=["petiscos","tábua de frios","porções quentes","pão de alho/mini-sanduíches"]
    elif req.type == "casamento": foods=["finger foods","canapés","jantar ou prato principal","doces finos","bolo"]
    elif req.type == "firma": foods=["finger foods","mini-sanduíches","doces/bolo","café e bebidas"]
    elif req.type.startswith("aniversario"):
        foods=["salgados","doces","bolo","mini-sanduíches ou lanches"]
    else: foods=["salgados/finger foods","bolo ou sobremesa","prato principal se houver refeição","frutas ou saladas"]
    sections=[
        {"title":"🍽️ Comidas", "items":[{"name":x,"quantity":"adaptar ao formato","reason":"Combina com o perfil do evento."} for x in foods]},
        {"title":"🥤 Bebidas", "items":[{"name":"Água","quantity":f"≈ {g*0.6:.1f} L","reason":"Base para todos os convidados."},{"name":"Refrigerantes/sucos","quantity":f"≈ {g*0.6:.1f} L","reason":"Complemento sem álcool."},{"name":"Bebida alcoólica","quantity":alcohol,"reason":"Somente para adultos e conforme o perfil informado."}]},
        {"title":"🍽️ Utensílios e infraestrutura", "items":[{"name":"Copos","quantity":f"≈ {g*4} un.","reason":"Margem para reposição."},{"name":"Pratos","quantity":f"≈ {g*2} un.","reason":"Aumente conforme o tipo de serviço."},{"name":"Talheres","quantity":f"≈ {max(g,round(g*1.5))} un.","reason":"Reserve extras."},{"name":"Cadeiras","quantity":f"≈ {g} un.","reason":"Ajustar ao formato e tempo em mesa."},{"name":"Guardanapos","quantity":f"≈ {g*2} un.","reason":"Consumo comum de evento."}]},
    ]
    return {"summary":f"Plano-base para {labels.get(req.type, req.type)} com {g} convidados.","sections":sections,"source":"fallback"}



@app.get("/api/address/cep/{cep}")
def address_by_cep(cep: str) -> dict[str, Any]:
    return _lookup_cep(cep)


@app.post("/api/address/search")
def address_search(req: SolarResourceRequest) -> dict[str, Any]:
    loc = geocode_solar(req)
    return {
        "lat": loc["lat"],
        "lon": loc["lon"],
        "displayName": loc.get("displayName", ""),
        "geocoder": loc.get("geocoder", "OpenStreetMap/Nominatim"),
        "address": loc.get("input", {}),
    }


@app.post("/api/solar/resource")
def solar_resource(req: SolarResourceRequest) -> dict[str, Any]:
    loc=geocode_solar(req)
    summary=pvgis_summary(loc["lat"],loc["lon"],float(req.roofAzimuth or 0),float(req.roofTilt or 15))
    return {**loc,**summary,"lat":loc["lat"],"lon":loc["lon"]}


@app.post("/api/solar/image-analyze")
def solar_image_analyze(req: SolarImageRequest) -> dict[str, Any]:
    if not OPENAI_API_KEY:
        raise HTTPException(status_code=503, detail="Configure OPENAI_API_KEY para habilitar análise visual por IA.")
    if not req.imageData.startswith("data:image/"):
        raise HTTPException(status_code=400, detail="A imagem precisa ser enviada como data URL.")
    schema={"type":"object","additionalProperties":False,"properties":{
        "summary":{"type":"string"},
        "confidence":{"type":"string"},
        "recommendedAzimuth":{"type":"number"},
        "zones":{"type":"array","items":{"type":"object","additionalProperties":False,"properties":{"x":{"type":"number"},"y":{"type":"number"},"width":{"type":"number"},"height":{"type":"number"},"reason":{"type":"string"}},"required":["x","y","width","height","reason"]}}
    },"required":["summary","confidence","recommendedAzimuth","zones"]}
    try:
        from openai import OpenAI
        client=OpenAI(api_key=OPENAI_API_KEY)
        prompt=f"""Analise a imagem de um imóvel/telhado para um estudo preliminar de posicionamento de módulos fotovoltaicos. A imagem pode ser um print de mapa/satélite e pode conter norte, ruas e outras áreas. Não assuma medidas exatas nem diga que um ponto é estruturalmente seguro. Identifique visualmente áreas de telhado que parecem utilizáveis, obstáculos visíveis e uma orientação recomendada de módulos. Retorne coordenadas normalizadas 0 a 1 para até 4 zonas retangulares aproximadas (centro x/y, largura/altura). A recomendação de azimute usa convenção 0=N, 90=L, 180=S, 270=O. Considere a localização informada apenas como contexto: lat={req.lat}, lon={req.lon}."""
        response=client.responses.create(model=OPENAI_MODEL,input=[{"role":"user","content":[{"type":"input_text","text":prompt},{"type":"input_image","image_url":req.imageData,"detail":"high"}]}],text={"format":{"type":"json_schema","name":"resolvei_solar_image","strict":True,"schema":schema}})
        return json.loads(response.output_text)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Falha na análise visual por IA: {exc}") from exc


@app.get("/api/solar/prices")
def solar_prices(city: str="Cuiabá",state: str="MT") -> dict[str, Any]:
    if not SERPAPI_KEY:
        return {"configured":False,"items":[]}
    queries={"panel":"painel solar fotovoltaico 550W","inverter":"inversor solar 5kW on grid","mounting":"estrutura solar telhado por módulo","dcProtection":"string box proteção DC fotovoltaica","acProtection":"quadro proteção AC fotovoltaico"}
    items=[]
    for key,q in queries.items():
        opts=search_price(q,city,state.upper())
        if opts:
            items.append({"key":key,"query":q,"best":opts[0],"options":opts})
    return {"configured":True,"city":city,"state":state.upper(),"items":items}


@app.post("/api/files/convert")
async def convert_file(file: UploadFile = File(...), output_format: str = Form(...)):
    allowed = {
        "mp4": {"avi","webm","mov"}, "avi": {"mp4","webm"}, "mov": {"mp4","avi","webm"},
        "mkv": {"mp4","webm"}, "webm": {"mp4","avi"},
        "jpg": {"png","webp","pdf"}, "jpeg": {"png","webp","pdf"},
        "png": {"jpg","webp","pdf"}, "webp": {"jpg","png","pdf"},
        "bmp": {"jpg","png","webp"}, "pdf": {"jpg","png","dxf","dwg"}
    }
    filename = Path(file.filename or "arquivo").name
    src_ext = Path(filename).suffix.lower().lstrip(".")
    out_ext = output_format.lower().lstrip(".")
    if src_ext not in allowed or out_ext not in allowed[src_ext]:
        raise HTTPException(status_code=400, detail="Conversão não suportada.")
    max_bytes = 200 * 1024 * 1024
    with tempfile.TemporaryDirectory(prefix="resolvei-convert-") as td:
        src = Path(td) / filename
        with src.open("wb") as f:
            total = 0
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk: break
                total += len(chunk)
                if total > max_bytes: raise HTTPException(status_code=413, detail="O arquivo excede o limite de 200 MB.")
                f.write(chunk)
        out = Path(td) / f"convertido.{out_ext}"
        try:
            if src_ext in {"jpg","jpeg","png","webp","bmp"}:
                from PIL import Image
                img = Image.open(src)
                if out_ext == "pdf":
                    img.convert("RGB").save(out, "PDF")
                else:
                    if out_ext == "jpg": out = out.with_suffix(".jpg"); img = img.convert("RGB")
                    img.save(out, format=out_ext.upper())
            elif src_ext == "pdf":
                import fitz
                doc = fitz.open(src)
                if len(doc) == 0: raise ValueError("PDF vazio.")
                if out_ext in {"jpg","png"}:
                    if len(doc) == 1:
                        pix = doc[0].get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
                        pix.save(str(out))
                    else:
                        out = out.with_suffix(".zip")
                        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
                            for i, page in enumerate(doc):
                                pix = page.get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
                                img = Path(td) / f"pagina-{i+1}.{out_ext}"
                                pix.save(str(img))
                                z.write(img, img.name)
                elif out_ext in {"dxf","dwg"}:
                    import ezdxf
                    cad = ezdxf.new("R2018")
                    msp = cad.modelspace()
                    for page_no, page in enumerate(doc):
                        layer=f"PDF_PAGE_{page_no+1}"
                        for item in page.get_drawings():
                            op=item[0]
                            if op=="l":
                                p1,p2=item[1],item[2]; msp.add_line((p1.x,-p1.y),(p2.x,-p2.y),dxfattribs={"layer":layer})
                            elif op=="re":
                                r=item[1]; pts=[(r.x0,-r.y0),(r.x1,-r.y0),(r.x1,-r.y1),(r.x0,-r.y1)]
                                msp.add_lwpolyline(pts,close=True,dxfattribs={"layer":layer})
                            elif op=="qu":
                                q=item[1]; pts=[(q.ul.x,-q.ul.y),(q.ur.x,-q.ur.y),(q.lr.x,-q.lr.y),(q.ll.x,-q.ll.y)]
                                msp.add_lwpolyline(pts,close=True,dxfattribs={"layer":layer})
                        for block in page.get_text("dict").get("blocks",[]):
                            for line in block.get("lines",[]):
                                for span in line.get("spans",[]):
                                    txt=span.get("text","").strip()
                                    if txt:
                                        x,y=span["origin"]; msp.add_text(txt,dxfattribs={"height":max(float(span.get("size",8)),1),"layer":f"TEXT_PAGE_{page_no+1}"}).set_placement((x,-y))
                    dxf_path=out.with_suffix(".dxf")
                    cad.saveas(dxf_path)
                    if out_ext=="dwg":
                        oda=shutil.which("ODAFileConverter") or shutil.which("odafileconverter")
                        if not oda:
                            raise RuntimeError("DWG exige ODA File Converter no servidor. Use DXF ou instale o conversor ODA.")
                        target=Path(td)/"dwgout"; target.mkdir()
                        p=subprocess.run([oda,str(td),str(target),"ACAD2018","DWG","0","1",str(dxf_path)],capture_output=True,text=True,timeout=300)
                        generated=list(target.rglob("*.dwg"))
                        if p.returncode!=0 or not generated: raise RuntimeError("Falha ao gerar DWG.")
                        shutil.copy2(generated[0],out)
                doc.close()
            else:
                ffmpeg = shutil.which("ffmpeg")
                if not ffmpeg: raise RuntimeError("FFmpeg não está instalado no servidor.")
                cmd = [ffmpeg, "-y", "-i", str(src)]
                if out_ext == "avi": cmd += ["-c:v","mpeg4","-c:a","mp3"]
                elif out_ext == "webm": cmd += ["-c:v","libvpx-vp9","-c:a","libopus"]
                elif out_ext == "mov": cmd += ["-c:v","libx264","-c:a","aac"]
                else: cmd += ["-c:v","libx264","-c:a","aac"]
                cmd += [str(out)]
                p = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
                if p.returncode != 0: raise RuntimeError("FFmpeg não conseguiu converter o vídeo.")
            media = {
                "jpg":"image/jpeg","png":"image/png","webp":"image/webp","pdf":"application/pdf","dxf":"application/dxf","dwg":"application/acad",
                "mp4":"video/mp4","avi":"video/x-msvideo","mov":"video/quicktime","webm":"video/webm"
            }.get(out_ext, "application/octet-stream")
            return Response(content=out.read_bytes(), media_type=media,
                            headers={"Content-Disposition": f'attachment; filename="resolvei-convertido.{("zip" if out_ext in {"jpg","png"} and src_ext=="pdf" and out.suffix==".zip" else out_ext)}"'});
        except HTTPException:
            raise
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=504, detail="A conversão demorou demais e foi interrompida.")
        except Exception as e:
            raise HTTPException(status_code=500, detail=str(e) or "Falha na conversão.")


@app.post("/api/files/convert-plus")
async def convert_plus(files: list[UploadFile] = File(...), output_format: str = Form(...), quality: int = Form(85), tool_id: str = Form(""), width: int = Form(0), height: int = Form(0)):
    """General-purpose document/media conversion endpoint used by the expanded Resolvei converters."""
    output_format = output_format.lower().lstrip(".")
    quality = max(10, min(100, int(quality or 85)))
    if not files:
        raise HTTPException(status_code=400, detail="Selecione pelo menos um arquivo.")
    max_bytes = 200 * 1024 * 1024
    with tempfile.TemporaryDirectory(prefix="resolvei-plus-") as td:
        paths=[]
        for up in files:
            name=Path(up.filename or "arquivo").name
            p=Path(td)/name
            total=0
            with p.open("wb") as f:
                while True:
                    chunk=await up.read(1024*1024)
                    if not chunk: break
                    total += len(chunk)
                    if total > max_bytes:
                        raise HTTPException(status_code=413, detail=f"{name} excede o limite de 200 MB.")
                    f.write(chunk)
            paths.append(p)
        try:
            if tool_id == "zip-arquivos":
                out=Path(td)/"resolvei-arquivos.zip"
                with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
                    for p in paths: z.write(p,p.name)
                return Response(content=out.read_bytes(),media_type="application/zip",headers={"Content-Disposition":'attachment; filename="resolvei-arquivos.zip"'})

            src=paths[0]
            ext=src.suffix.lower().lstrip(".")
            out=Path(td)/f"resolvei-convertido.{output_format}"

            if tool_id in {"jpg-png-webp","imagem-comprimir","heic-jpg","jpg-heic"}:
                from PIL import Image
                img=Image.open(src)
                if output_format in {"jpg","jpeg"}:
                    out=out.with_suffix(".jpg")
                    img.convert("RGB").save(out,"JPEG",quality=quality,optimize=True)
                elif output_format=="png":
                    img.save(out,"PNG",optimize=True)
                elif output_format=="webp":
                    img.save(out,"WEBP",quality=quality,method=6)
                elif output_format=="heic":
                    try:
                        from pillow_heif import register_heif_opener
                        register_heif_opener()
                        img.convert("RGB").save(out,"HEIC",quality=quality)
                    except Exception as exc:
                        raise RuntimeError("HEIC não está disponível neste servidor. Instale pillow-heif para habilitar.") from exc
            elif tool_id=="imagem-pdf":
                from PIL import Image
                images=[]
                for p in paths:
                    im=Image.open(p).convert("RGB")
                    images.append(im)
                if not images: raise ValueError("Nenhuma imagem válida.")
                images[0].save(out,"PDF",save_all=True,append_images=images[1:])
            elif tool_id=="pdf-imagens-zip":
                if ext!="pdf": raise ValueError("Envie um PDF.")
                import fitz
                doc=fitz.open(src)
                out=Path(td)/"resolvei-pdf-imagens.zip"
                with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
                    for i,page in enumerate(doc):
                        pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False)
                        img=Path(td)/f"pagina-{i+1}.jpg"
                        pix.save(str(img)); z.write(img,img.name)
                doc.close()
            elif tool_id=="csv-xlsx":
                if ext=="csv" and output_format=="xlsx":
                    import csv, openpyxl
                    wb=openpyxl.Workbook(); ws=wb.active
                    with src.open("r",encoding="utf-8-sig",newline="") as f:
                        for row in csv.reader(f): ws.append(row)
                    wb.save(out)
                elif ext=="xlsx" and output_format=="csv":
                    import openpyxl, csv
                    wb=openpyxl.load_workbook(src,read_only=True,data_only=True); ws=wb.active
                    with out.open("w",encoding="utf-8-sig",newline="") as f:
                        w=csv.writer(f)
                        for row in ws.iter_rows(values_only=True): w.writerow(list(row))
                else: raise ValueError("Conversão CSV/XLSX inválida.")
            elif tool_id in {"mp4-mp3"}:
                ffmpeg=shutil.which("ffmpeg")
                if not ffmpeg: raise RuntimeError("FFmpeg não está instalado no servidor.")
                out=out.with_suffix(".mp3")
                p=subprocess.run([ffmpeg,"-y","-i",str(src),"-vn","-codec:a","libmp3lame","-q:a","2",str(out)],capture_output=True,text=True,timeout=600)
                if p.returncode!=0: raise RuntimeError("FFmpeg não conseguiu extrair o áudio.")
            elif tool_id in {"mp4-gif","mov-mp4","video-webm","video-avi","video-audio"}:
                ffmpeg=shutil.which("ffmpeg")
                if not ffmpeg: raise RuntimeError("FFmpeg não está instalado no servidor.")
                if tool_id=="mp4-gif":
                    out=out.with_suffix(".gif"); cmd=[ffmpeg,"-y","-i",str(src),"-vf","fps=12,scale=640:-1:flags=lanczos","-t","10",str(out)]
                elif tool_id=="video-webm":
                    out=out.with_suffix(".webm"); cmd=[ffmpeg,"-y","-i",str(src),"-c:v","libvpx-vp9","-c:a","libopus",str(out)]
                elif tool_id=="video-avi":
                    out=out.with_suffix(".avi"); cmd=[ffmpeg,"-y","-i",str(src),"-c:v","mpeg4","-c:a","mp3",str(out)]
                elif tool_id=="video-audio":
                    out=out.with_suffix(".mp3"); cmd=[ffmpeg,"-y","-i",str(src),"-vn","-codec:a","libmp3lame","-q:a","2",str(out)]
                else:
                    out=out.with_suffix(".mp4"); cmd=[ffmpeg,"-y","-i",str(src),"-c:v","libx264","-c:a","aac","-movflags","+faststart",str(out)]
                p=subprocess.run(cmd,capture_output=True,text=True,timeout=600)
                if p.returncode!=0: raise RuntimeError("FFmpeg não conseguiu converter o vídeo.")
            elif tool_id in {"audio-mp3-wav","audio-ogg"}:
                ffmpeg=shutil.which("ffmpeg")
                if not ffmpeg: raise RuntimeError("FFmpeg não está instalado no servidor.")
                out=out.with_suffix("."+output_format)
                codec="pcm_s16le" if output_format=="wav" else "libvorbis"
                p=subprocess.run([ffmpeg,"-y","-i",str(src),"-vn","-codec:a",codec,str(out)],capture_output=True,text=True,timeout=600)
                if p.returncode!=0: raise RuntimeError("FFmpeg não conseguiu converter o áudio.")
            elif tool_id=="docx-pdf":
                from docx import Document
                from reportlab.lib.pagesizes import A4
                from reportlab.pdfgen import canvas
                doc=Document(src); out=out.with_suffix(".pdf"); pdf=canvas.Canvas(str(out),pagesize=A4); w,h=A4; y=h-50
                for para in doc.paragraphs:
                    text=para.text.strip()
                    if not text: y-=12; continue
                    for line in text.splitlines():
                        pdf.drawString(40,y,line[:120]); y-=14
                        if y<45: pdf.showPage(); y=h-50
                pdf.save()
            elif tool_id=="pdf-docx":
                from docx import Document
                import fitz
                docx=Document(); pdf=fitz.open(src)
                for i,page in enumerate(pdf):
                    if i: docx.add_page_break()
                    for line in page.get_text("text").splitlines(): docx.add_paragraph(line)
                out=out.with_suffix(".docx"); docx.save(out); pdf.close()
            elif tool_id=="pdf-txt":
                import fitz
                pdf=fitz.open(src); out=out.with_suffix(".txt"); out.write_text("\n\n".join(p.get_text("text") for p in pdf),encoding="utf-8"); pdf.close()
            elif tool_id=="txt-pdf":
                from reportlab.lib.pagesizes import A4
                from reportlab.pdfgen import canvas
                out=out.with_suffix(".pdf"); pdf=canvas.Canvas(str(out),pagesize=A4); w,h=A4; y=h-50
                for line in src.read_text(encoding="utf-8",errors="replace").splitlines():
                    pdf.drawString(40,y,line[:120]); y-=14
                    if y<45: pdf.showPage(); y=h-50
                pdf.save()
            elif tool_id=="pdf-xlsx":
                import fitz, openpyxl
                pdf=fitz.open(src); wb=openpyxl.Workbook(); ws=wb.active; ws.title="PDF"
                for pageno,page in enumerate(pdf,1):
                    for line in page.get_text("text").splitlines(): ws.append([pageno,line])
                out=out.with_suffix(".xlsx"); wb.save(out); pdf.close()
            elif tool_id=="xlsx-csv":
                import openpyxl, csv
                wb=openpyxl.load_workbook(src,read_only=True,data_only=True); ws=wb.active; out=out.with_suffix(".csv")
                with out.open("w",encoding="utf-8-sig",newline="") as f:
                    w=csv.writer(f)
                    for row in ws.iter_rows(values_only=True): w.writerow(list(row))
            elif tool_id=="svg-png":
                try:
                    import cairosvg
                    out=out.with_suffix(".png"); cairosvg.svg2png(url=str(src),write_to=str(out),output_width=width or None,output_height=height or None)
                except Exception as exc: raise RuntimeError("SVG→PNG requer cairosvg no servidor.") from exc
            elif tool_id=="png-ico":
                from PIL import Image
                img=Image.open(src).convert("RGBA"); out=out.with_suffix(".ico"); img.save(out,format="ICO",sizes=[(16,16),(32,32),(48,48),(256,256)])
            elif tool_id=="imagem-redimensionar":
                from PIL import Image
                img=Image.open(src); w=width or img.width; h=height or img.height; img=img.resize((w,h),Image.Resampling.LANCZOS); out=out.with_suffix("."+output_format)
                if output_format=="jpg": img=img.convert("RGB")
                img.save(out,format=output_format.upper(),quality=quality if output_format in {"jpg","webp"} else None)
            elif tool_id=="pdf-comprimir":
                import fitz
                pdf=fitz.open(src); out=out.with_suffix(".pdf"); pdf.save(out,garbage=4,deflate=True,clean=True); pdf.close()
            elif tool_id=="arquivos-zip":
                out=Path(td)/"resolvei-arquivos.zip"
                with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
                    for p in paths: z.write(p,p.name)
            else:
                raise ValueError("Conversor não reconhecido.")

            mime={"jpg":"image/jpeg","png":"image/png","webp":"image/webp","heic":"image/heic","pdf":"application/pdf","zip":"application/zip","xlsx":"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet","csv":"text/csv","mp3":"audio/mpeg","gif":"image/gif","mp4":"video/mp4"}.get(output_format,"application/octet-stream")
            return Response(content=out.read_bytes(),media_type=mime,headers={"Content-Disposition":f'attachment; filename="{out.name}"'})
        except HTTPException:
            raise
        except subprocess.TimeoutExpired:
            raise HTTPException(status_code=504,detail="A conversão demorou demais e foi interrompida.")
        except Exception as exc:
            raise HTTPException(status_code=500,detail=str(exc) or "Falha na conversão.")

@app.get("/api/health")
def health() -> dict[str, Any]:
    return {"ok": True, "ai_configured": bool(OPENAI_API_KEY or GEMINI_API_KEY), "gemini_configured": bool(GEMINI_API_KEY), "prices_configured": bool(SERPAPI_KEY), "firebase_admin_configured": bool(_firebase_admin), "encryption_configured": bool(RESOLVEI_CREDENTIAL_ENCRYPTION_KEY), "model": OPENAI_MODEL, "gemini_model": GEMINI_MODEL}


@app.post("/api/recipe/analyze")
def analyze_recipe(req: RecipeRequest, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    url = safe_url(req.url)
    html = fetch_html(url)
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else "Receita"
    raw = extract_raw_ingredients(html)
    if not raw:
        raise HTTPException(status_code=422, detail="Não encontrei ingredientes estruturados nessa página. Tente outra URL ou insira os ingredientes manualmente.")
    ingredients = normalize_ingredients(raw, title)
    # Usuários autenticados podem usar uma conexão pessoal ou o Gemini do Resolvei.
    # A chave nunca é devolvida ao navegador.
    price_api_key=req.api_key.strip()
    price_provider=req.provider.strip().lower()
    price_model=req.model.strip()
    if not price_api_key and authorization:
        try:
            price_provider,price_api_key,stored_model=_resolve_ai_credentials(authorization,price_provider)
            price_model=price_model or stored_model
        except HTTPException:
            price_api_key=""
    ai_price_error = ""
    ai_price_count = 0
    if price_api_key:
        price_prompt = f"""Você é um estimador de preços de supermercado no Brasil. Analise os ingredientes abaixo para a receita "{title}". Para cada ingrediente, estime o preço em reais da quantidade efetivamente usada na receita, considerando preços típicos e atuais para {req.city}, {req.state.upper()}. Não confunda o preço da embalagem inteira com o custo proporcional à quantidade usada. Seja conservador quando houver variação regional. Retorne somente JSON válido no formato {{"items":[{{"name":"...","price":0.0,"source":"estimativa IA"}}]}}. Inclua todos os ingredientes. Nunca invente links ou lojas. Ingredientes: """ + json.dumps([{"name":x.get("name",""),"quantity":x.get("quantity",1),"unit":x.get("unit","un.")} for x in ingredients], ensure_ascii=False)
        try:
            raw_prices = _provider_call(price_provider or "gemini", price_api_key, price_model, price_prompt)
            raw_prices = re.sub(r"^\`\`\`(?:json)?\s*|\s*\`\`\`$", "", raw_prices.strip(), flags=re.I)
            parsed = json.loads(raw_prices)
            estimated = parsed.get("items", []) if isinstance(parsed, dict) else []
            by_name={str(x.get("name","")).strip().lower():x for x in estimated if isinstance(x,dict)}
            for ing in ingredients:
                p=by_name.get(str(ing.get("name","")).strip().lower())
                if p and isinstance(p.get("price"),(int,float)):
                    ing["recipe_price"]=float(p["price"])
                    ai_price_count += 1
            if ai_price_count == 0:
                ai_price_error = "A IA respondeu, mas não retornou preços em formato válido."
        except Exception as exc:
            ai_price_error = f"Não foi possível estimar os preços pela IA: {str(exc)[:180]}"
    price_items: list[dict[str, Any]] = []
    basket_total = 0.0
    if SERPAPI_KEY:
        for ing in ingredients[:12]:
            options = search_price(ing.get("name", ""), req.city or "Cuiabá", (req.state or "MT").upper())
            best = options[0] if options else None
            price_items.append({"name": ing.get("name", ""), "price": best.get("price") if best else None, "source": best.get("source") if best else "Sem cotação", "title": best.get("title") if best else "", "options": options})
            if best:
                basket_total += float(best["price"])
        price_note = f"Cotação de referência para até {min(12, len(ingredients))} ingredientes, orientada por {req.city}, {req.state}. A soma representa produtos encontrados, não o custo proporcional efetivamente consumido na receita."
    else:
        price_note = "Os ingredientes foram organizados, mas a cotação regional ao vivo está desativada. Configure SERPAPI_KEY no servidor para consultar preços."
    return {
        "title": title,
        "ingredients": ingredients,
        "price_summary": {"total": basket_total, "items": price_items} if SERPAPI_KEY else None,
        "price_note": price_note,
        "ai_price_count": ai_price_count,
        "ai_price_error": ai_price_error,
        "city": req.city,
        "state": req.state,
    }


@app.post("/api/sales/price")
def sales_price(req: SalesPriceRequest) -> dict[str, Any]:
    # Precificação determinística: sem IA, sem API key e sem dados externos.
    if not req.product.strip():
        raise HTTPException(status_code=400, detail="Informe o produto que você vende.")
    qty=max(float(req.quantity or 0),1.0)
    production=max(float(req.production_cost or 0),0.0)
    packaging=max(float(req.packaging_per_unit or 0),0.0)
    other=max(float(req.other_costs or 0),0.0)
    # Campos adicionais são enviados pelo frontend através dos valores já incorporados
    # no custo de produção/outros gastos. O endpoint mantém compatibilidade com versões anteriores.
    total_batch=production + (packaging*qty) + other
    unit_cost=total_batch/qty
    if unit_cost<=0:
        raise HTTPException(status_code=400, detail="Informe pelo menos o custo dos ingredientes/preparo ou outro gasto do lote.")
    fees=min(max(float(getattr(req,"fees_percent",0) or 0),0.0),100.0)
    tax=min(max(float(getattr(req,"tax_percent",0) or 0),0.0),100.0)
    margin=min(max(float(getattr(req,"target_margin",30) or 30),0.0),90.0)
    rate=(fees+tax)/100.0
    if rate>=1:
        raise HTTPException(status_code=400, detail="Taxas e impostos precisam ficar abaixo de 100%.")
    if rate + margin/100 >= 1:
        raise HTTPException(status_code=400, detail="A soma de taxas, impostos e lucro desejado ultrapassa 100%. Reduza a margem.")
    break_even=unit_cost/(1-rate)
    suggested=unit_cost/(1-rate-margin/100)
    price20=unit_cost/(1-rate-0.20) if rate<0.80 else None
    price40=unit_cost/(1-rate-0.40) if rate<0.60 else None
    net_per_unit=suggested*(1-rate)-unit_cost
    return {
        "product": req.product.strip(),
        "city": (req.city or "").strip(),
        "state": (req.state or "").strip().upper(),
        "quantity": qty,
        "total_batch_cost": total_batch,
        "unit_cost": unit_cost,
        "break_even_price": break_even,
        "minimum_price": break_even,
        "suggested_price": suggested,
        "price_20_margin": price20,
        "price_40_margin": price40,
        "profit_per_unit": max(net_per_unit,0.0),
        "batch_profit": max(net_per_unit,0.0)*qty,
        "fees_percent": fees,
        "tax_percent": tax,
        "target_margin": margin,
        "note": "Cálculo matemático baseado nos custos informados. A cidade/UF serve para contextualizar a comparação com o mercado local; não é uma cotação automática de preços."
    }

@app.post("/api/party/suggest")
def party_suggest(req: PartyRequest, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    req.drinkers = min(req.drinkers, req.adults + req.kids)
    schema = {
        "type":"object","additionalProperties":False,
        "properties":{
            "summary":{"type":"string"},
            "sections":{"type":"array","items":{"type":"object","additionalProperties":False,"properties":{
                "title":{"type":"string"},
                "items":{"type":"array","items":{"type":"object","additionalProperties":False,"properties":{"name":{"type":"string"},"quantity":{"type":"string"},"reason":{"type":"string"}},"required":["name","quantity","reason"]}}
            },"required":["title","items"]}}
        },"required":["summary","sections"]
    }
    labels={"aniversario-infantil":"aniversário infantil","aniversario-adulto":"aniversário adulto","casamento":"casamento","firma":"festa da firma","cha-bebe":"chá de bebê","cha-revelacao":"chá revelação","noivado":"noivado","formatura":"formatura","bodas":"bodas","familiar":"comemoração familiar","junina":"festa junina","happy-hour":"happy hour","outro":"outra comemoração"}
    prompt=f"""Você é um planejador de festas para o Brasil. Monte um plano prático para: tipo={labels.get(req.type,req.type)}, idade do aniversariante={req.age}, adultos={req.adults}, crianças={req.kids}, duração={req.hours} horas, bebedores de álcool={req.drinkers}, bebida={req.alcohol}, orçamento por pessoa={req.budgetPerPerson}. Sugira comidas, bebidas alcoólicas e não alcoólicas, drinks quando apropriado, gelo, utensílios, copos, pratos, talheres, guardanapos, cadeiras e infraestrutura. Adapte o foco ao tipo e, em aniversário, considere a idade. Não invente preços. Dê quantidades em linguagem prática e explique brevemente o motivo de cada item. Para menores, não sugira consumo de álcool."""
    if authorization:
        try:
            provider,api_key,model=_resolve_ai_credentials(authorization,"")
            raw=_provider_call(provider,api_key,model,prompt)
            match=re.search(r"\{.*\}",raw,re.S)
            if match:
                data=json.loads(match.group(0))
                if isinstance(data,dict):
                    data["source"]=provider
                    return data
        except Exception:
            pass
    if OPENAI_API_KEY:
        try:
            data=openai_json(prompt,"resolvei_party",schema)
            data["source"]="openai"
            return data
        except RuntimeError:
            pass
    result=fallback_party(req)
    result["ai_note"]="A IA não está configurada/indisponível; o Resolvei usou o plano-base inteligente." if not OPENAI_API_KEY else "A IA ficou indisponível; o Resolvei usou o plano-base inteligente."
    return result


@app.get("/api/risk/state")
def risk_state_get(authorization: str | None = Header(default=None)) -> dict[str, Any]:
    user = _verify_firebase_token(authorization)
    if not _firebase_admin:
        raise HTTPException(status_code=503, detail="Firebase indisponível.")
    from firebase_admin import firestore as firebase_firestore
    ref = (
        firebase_firestore.client()
        .collection("users").document(user["uid"])
        .collection("riskProgress").document("capitalEvolution")
    )
    snap = ref.get()
    if not snap.exists:
        return {"exists": False, "state": None}
    data = snap.to_dict() or {}
    state = data.get("state")
    return {"exists": isinstance(state, dict), "state": state if isinstance(state, dict) else None}


@app.put("/api/risk/state")
def risk_state_save(req: RiskProgressRequest, authorization: str | None = Header(default=None)) -> dict[str, Any]:
    user = _verify_firebase_token(authorization)
    if not _firebase_admin:
        raise HTTPException(status_code=503, detail="Firebase indisponível.")
    state = req.state if isinstance(req.state, dict) else {}
    try:
        encoded = json.dumps(state, ensure_ascii=False, separators=(",", ":"))
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="Estado do plano inválido.") from exc
    if len(encoded) > 500_000:
        raise HTTPException(status_code=413, detail="O histórico do plano está muito grande.")
    from firebase_admin import firestore as firebase_firestore
    ref = (
        firebase_firestore.client()
        .collection("users").document(user["uid"])
        .collection("riskProgress").document("capitalEvolution")
    )
    ref.set({
        "state": state,
        "schemaVersion": 1,
        "updatedAt": firebase_firestore.SERVER_TIMESTAMP
    }, merge=True)
    return {"ok": True, "saved": True}


@app.post("/api/prices/search")
def prices_search(req: ShoppingPriceRequest) -> dict[str, Any]:
    item=req.item.strip()
    if not item:
        raise HTTPException(status_code=400, detail="Informe um item.")
    options=search_price(item,req.city,req.state.upper())
    return {"item":item,"city":req.city,"state":req.state.upper(),"options":options}


@app.get("/{path:path}")
def static_files(path: str = ""):
    # Serve normal files; for browser routes return index.html.
    target=(BASE_DIR/path).resolve()
    try:
        target.relative_to(BASE_DIR.resolve())
    except ValueError:
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    if path and target.is_file():
        return FileResponse(target)
    index=BASE_DIR/"index.html"
    return FileResponse(index)
