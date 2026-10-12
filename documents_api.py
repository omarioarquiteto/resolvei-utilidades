from __future__ import annotations

import difflib
import io
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import fitz
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import Response

router = APIRouter(prefix="/api/documents", tags=["documentos-profissionais"])
MAX_FILE_BYTES = 200 * 1024 * 1024
MAX_PDF_PAGES = 100


async def _save_upload(upload: UploadFile, target: Path) -> int:
    total = 0
    with target.open("wb") as stream:
        while True:
            chunk = await upload.read(1024 * 1024)
            if not chunk:
                break
            total += len(chunk)
            if total > MAX_FILE_BYTES:
                raise HTTPException(status_code=413, detail=f"{Path(upload.filename or 'arquivo').name} excede 200 MB.")
            stream.write(chunk)
    return total


def _read_text(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        try:
            with fitz.open(path) as doc:
                if len(doc) > MAX_PDF_PAGES:
                    raise HTTPException(status_code=413, detail="O PDF ultrapassa o limite de 100 páginas.")
                return [page.get_text("text", sort=True) for page in doc]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Não foi possível ler o PDF. Verifique se o arquivo está íntegro e não possui senha.") from exc
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return ["\n".join(parts)]
        except Exception as exc:
            raise HTTPException(status_code=400, detail="Não foi possível ler o DOCX.") from exc
    if ext in {".txt", ".csv"}:
        try:
            return [path.read_text(encoding="utf-8-sig")]
        except UnicodeDecodeError:
            return [path.read_text(encoding="latin-1")]
    raise HTTPException(status_code=415, detail="Formato não suportado. Use PDF, DOCX ou TXT.")


def _ocr_image(image: Any) -> str:
    """Use Tesseract when available and a Python/ONNX OCR engine as fallback."""
    try:
        import pytesseract
        if shutil.which("tesseract"):
            text = pytesseract.image_to_string(image, lang="por+eng").strip()
            if text:
                return text
    except Exception:
        pass
    try:
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
        engine = _rapid_ocr_engine()
        result, _elapsed = engine(np.asarray(image))
        return "\\n".join(str(item[1]).strip() for item in (result or []) if len(item) > 1 and str(item[1]).strip())
    except Exception as exc:
        raise RuntimeError("Nenhum motor OCR disponível (Tesseract ou RapidOCR).") from exc


_rapid_engine = None


def _rapid_ocr_engine():
    global _rapid_engine
    if _rapid_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _rapid_engine = RapidOCR()
    return _rapid_engine


def _download(data: bytes, filename: str, media_type: str) -> Response:
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", filename)
    return Response(content=data, media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{safe}"', "Cache-Control": "no-store"})


@router.post("/pdfa")
async def convert_pdfa(file: UploadFile = File(...)) -> Response:
    filename = Path(file.filename or "documento.pdf").name
    if Path(filename).suffix.lower() != ".pdf":
        raise HTTPException(status_code=415, detail="Envie um arquivo PDF.")
    gs = shutil.which("gs") or shutil.which("gswin64c") or shutil.which("gswin32c")
    if not gs:
        raise HTTPException(status_code=503, detail="O servidor ainda não possui Ghostscript instalado. A conversão PDF/A fica indisponível até habilitar essa dependência no ambiente Render.")
    with tempfile.TemporaryDirectory(prefix="resolvei-pdfa-") as td:
        root = Path(td)
        src, out = root / "entrada.pdf", root / "documento-pdfa.pdf"
        await _save_upload(file, src)
        try:
            with fitz.open(src) as doc:
                if not len(doc):
                    raise HTTPException(status_code=400, detail="O PDF está vazio.")
                if doc.needs_pass:
                    raise HTTPException(status_code=400, detail="PDF protegido por senha não é suportado.")
            icc_candidates = [
                Path("/usr/share/color/icc/ghostscript/srgb.icc"),
                Path("/usr/share/color/icc/colord/sRGB.icc"),
                Path("/usr/share/ghostscript/iccprofiles/default_rgb.icc"),
            ]
            icc = next((p for p in icc_candidates if p.exists()), None)
            if not icc:
                # Gera um perfil ICC sRGB local com Pillow, evitando depender de um
                # caminho específico de distribuição Linux.
                try:
                    from PIL import ImageCms
                    profile = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB"))
                    icc = root / "sRGB.icc"
                    icc.write_bytes(profile.tobytes())
                except Exception as exc:
                    raise HTTPException(status_code=503, detail="Não foi possível criar o perfil ICC sRGB exigido pelo PDF/A.") from exc
            definition = root / "PDFA_def.ps"
            definition.write_text(
                f"/ICCProfile ({str(icc)}) def\n"
                "[/_objdef {icc_PDFA} /type /stream /OBJ pdfmark\n"
                "[{icc_PDFA} << /N 3 >> /PUT pdfmark\n"
                "[{icc_PDFA} ICCProfile (r) file /PUT pdfmark\n"
                "[/_objdef {OutputIntent_PDFA} /type /dict /OBJ pdfmark\n"
                "[{OutputIntent_PDFA} << /Type /OutputIntent /S /GTS_PDFA1 "
                "/DestOutputProfile {icc_PDFA} /OutputConditionIdentifier (sRGB IEC61966-2.1) "
                "/Info (sRGB IEC61966-2.1) >> /PUT pdfmark\n"
                "[{Catalog} << /OutputIntents [ {OutputIntent_PDFA} ] >> /PUT pdfmark\n",
                encoding="latin-1",
            )
            cmd = [gs, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE", "-sDEVICE=pdfwrite",
                   "-dPDFACompatibilityPolicy=1", "-sColorConversionStrategy=RGB",
                   "-sProcessColorModel=DeviceRGB", f"-sOutputFile={out}", str(definition), str(src)]
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
            if proc.returncode != 0 or not out.exists() or out.stat().st_size == 0:
                detail = (proc.stderr or proc.stdout or "Ghostscript não conseguiu gerar o PDF/A.")[-1200:]
                raise HTTPException(status_code=422, detail=f"Falha na conversão PDF/A: {detail}")
            with fitz.open(out) as result:
                if not len(result):
                    raise HTTPException(status_code=422, detail="O arquivo convertido não passou na verificação estrutural.")
                metadata = result.metadata or {}
                basic_check = {"pages": len(result), "opens": True, "metadata_present": bool(metadata)}
            # Verificação estrutural é útil, mas não substitui uma validação formal por veraPDF.
            data = out.read_bytes()
            return _download(data, Path(filename).stem + "-PDF-A.pdf", "application/pdf")
        except HTTPException:
            raise
        except subprocess.TimeoutExpired as exc:
            raise HTTPException(status_code=504, detail="A conversão PDF/A excedeu o tempo limite.") from exc
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Não foi possível converter o PDF: {str(exc)[:500]}") from exc


@router.post("/ocr")
async def extract_ocr(file: UploadFile = File(...), output_format: str = Form("txt")) -> Response:
    filename = Path(file.filename or "documento").name
    ext = Path(filename).suffix.lower()
    if ext not in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}:
        raise HTTPException(status_code=415, detail="Use PDF ou imagem (PNG, JPG, TIFF, WEBP ou BMP).")
    if output_format.lower() != "txt":
        raise HTTPException(status_code=400, detail="Por enquanto, o OCR exporta texto TXT.")
    with tempfile.TemporaryDirectory(prefix="resolvei-ocr-") as td:
        src = Path(td) / filename
        await _save_upload(file, src)
        pages: list[str] = []
        try:
            if ext == ".pdf":
                with fitz.open(src) as doc:
                    if doc.needs_pass:
                        raise HTTPException(status_code=400, detail="PDF protegido por senha não é suportado.")
                    if len(doc) > MAX_PDF_PAGES:
                        raise HTTPException(status_code=413, detail="O PDF ultrapassa o limite de 100 páginas.")
                    pages = [p.get_text("text", sort=True).strip() for p in doc]
                    missing = [i for i, text in enumerate(pages) if len(text) < 15]
                    if missing:
                        try:
                            from PIL import Image
                            for i in missing:
                                pix = doc[i].get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                                image = Image.open(io.BytesIO(pix.tobytes("png")))
                                pages[i] = _ocr_image(image)
                        except Exception as exc:
                            if not any(pages):
                                raise HTTPException(status_code=503, detail="Este documento parece digitalizado, mas nenhum motor OCR conseguiu processá-lo. Tente uma imagem mais nítida ou reporte o problema.") from exc
                text = "\n\n".join(f"--- Página {i + 1} ---\n{content}" for i, content in enumerate(pages))
            else:
                try:
                    from PIL import Image, ImageOps
                    image = ImageOps.exif_transpose(Image.open(src))
                    text = _ocr_image(image)
                except Exception as exc:
                    raise HTTPException(status_code=503, detail="OCR indisponível. Verifique a imagem e tente novamente.") from exc
            if not text.strip():
                raise HTTPException(status_code=422, detail="Nenhum texto foi reconhecido. Tente uma imagem mais nítida ou verifique o idioma do documento.")
            return _download(text.encode("utf-8"), Path(filename).stem + "-ocr.txt", "text/plain; charset=utf-8")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Falha no OCR: {str(exc)[:400]}") from exc


@router.post("/tables")
async def extract_tables(file: UploadFile = File(...)) -> Response:
    filename = Path(file.filename or "documento.pdf").name
    if Path(filename).suffix.lower() != ".pdf":
        raise HTTPException(status_code=415, detail="Envie um PDF com tabelas.")
    with tempfile.TemporaryDirectory(prefix="resolvei-tables-") as td:
        src, out = Path(td) / filename, Path(td) / "tabelas.xlsx"
        await _save_upload(file, src)
        try:
            import pdfplumber
            from openpyxl import Workbook
            wb = Workbook()
            summary = wb.active
            summary.title = "Tabelas"
            found = 0
            with pdfplumber.open(src) as pdf:
                if len(pdf.pages) > MAX_PDF_PAGES:
                    raise HTTPException(status_code=413, detail="O PDF ultrapassa o limite de 100 páginas.")
                for page_no, page in enumerate(pdf.pages, start=1):
                    for table_no, table in enumerate(page.extract_tables() or [], start=1):
                        rows = [[str(cell or "").strip() for cell in row] for row in table if row and any(str(cell or "").strip() for cell in row)]
                        if not rows:
                            continue
                        if found:
                            summary.append([])
                        summary.append([f"Página {page_no} · Tabela {table_no}"])
                        for row in rows:
                            summary.append(row)
                        found += 1
            if found == 0:
                # Fallback transparente: exporta linhas de texto para uma folha separada,
                # sem alegar que conseguiu reconstruir uma tabela.
                wb.remove(summary)
                sheet = wb.create_sheet("Texto extraído")
                with fitz.open(src) as doc:
                    for page_no, page in enumerate(doc, start=1):
                        sheet.append([f"Página {page_no}"])
                        for line in page.get_text("text", sort=True).splitlines():
                            if line.strip():
                                sheet.append([line.strip()])
                if sheet.max_row <= len(doc):
                    raise HTTPException(status_code=422, detail="Não encontrei tabelas nem texto extraível. Se o PDF for digitalizado, use primeiro a ferramenta OCR.")
                sheet.insert_rows(1)
                sheet["A1"] = "Nenhuma tabela detectada: conteúdo exportado como linhas de texto para revisão."
            wb.save(out)
            return _download(out.read_bytes(), Path(filename).stem + "-tabelas.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=422, detail=f"Não foi possível extrair tabelas: {str(exc)[:500]}") from exc


@router.post("/compare")
async def compare_documents(files: list[UploadFile] = File(...)) -> Response:
    if len(files) != 2:
        raise HTTPException(status_code=400, detail="Selecione exatamente dois documentos para comparar.")
    with tempfile.TemporaryDirectory(prefix="resolvei-compare-") as td:
        paths = []
        for i, upload in enumerate(files, start=1):
            name = Path(upload.filename or f"documento-{i}.pdf").name
            target = Path(td) / f"{i}-{name}"
            await _save_upload(upload, target)
            paths.append(target)
        left_ext, right_ext = paths[0].suffix.lower(), paths[1].suffix.lower()
        supported = {".pdf", ".docx", ".txt", ".csv"}
        if left_ext not in supported or right_ext not in supported:
            raise HTTPException(status_code=415, detail="Compare dois arquivos PDF, DOCX, TXT ou CSV.")
        left_pages, right_pages = _read_text(paths[0]), _read_text(paths[1])
        left = "\n".join(left_pages).splitlines()
        right = "\n".join(right_pages).splitlines()
        diff = list(difflib.unified_diff(left, right, fromfile=Path(files[0].filename or "documento-1").name,
                                         tofile=Path(files[1].filename or "documento-2").name, lineterm=""))
        if not diff:
            report = "Os textos extraídos dos dois documentos são idênticos.\n\nNota: a comparação considera texto extraível, não diferenças visuais, assinaturas ou conteúdo em imagens."
        else:
            report = "RELATÓRIO DE COMPARAÇÃO — RESOLVEI\n\nLinhas com '-' existem apenas no primeiro arquivo; linhas com '+' existem apenas no segundo.\n\n" + "\n".join(diff)
            report += "\n\nNota: a comparação considera texto extraível, não diferenças visuais, assinaturas ou conteúdo em imagens."
        return _download(report.encode("utf-8"), "comparacao-documentos.txt", "text/plain; charset=utf-8")
